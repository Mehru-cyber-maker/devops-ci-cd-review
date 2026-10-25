"""Build a styled DOCX and PDF from the research manuscript. Requires pandoc."""

from pathlib import Path
import json
import re
import subprocess
import html
import zipfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    Preformatted, KeepTogether, Flowable, Image as ReportlabImage,
)
from reportlab.pdfgen import canvas as pdfcanvas

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / ".build"
BUILD.mkdir(exist_ok=True)
TABLE_RATIOS = {1:[.18,.26,.27,.29],2:[.11,.28,.61],3:[.21,.19,.30,.30],4:[.08,.18,.09,.09,.09,.09,.14,.14],5:[.16,.13,.14,.14,.16,.27],6:[.20,.09,.15,.56],7:[.22,.31,.17,.30]}
BLUE = '#18354d'
GRAY = '#52616b'


def make_reference():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.8)
    sec.left_margin = sec.right_margin = Inches(.8)
    sec.page_width, sec.page_height = Inches(8.2677), Inches(11.6929)
    styles = doc.styles
    for name in ['Normal', 'Body Text', 'First Paragraph']:
        if name not in styles:
            continue
        s = styles[name]
        s.font.name = 'Times New Roman'
        s.font.size = Pt(11)
        s.paragraph_format.line_spacing = 1.16
        s.paragraph_format.space_after = Pt(7)
    styles['Title'].font.name = 'Calibri'
    styles['Title'].font.size = Pt(25)
    styles['Title'].font.bold = True
    styles['Title'].font.color.rgb = RGBColor.from_string('18354D')
    styles['Subtitle'].font.name = 'Calibri'
    styles['Subtitle'].font.size = Pt(12)
    styles['Subtitle'].font.color.rgb = RGBColor.from_string('52616B')
    for name, size in [('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 11.5)]:
        s = styles[name]
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string('18354D')
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(6)
    if 'Source Code' in styles:
        styles['Source Code'].font.name = 'Consolas'
        styles['Source Code'].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run('EvidenceGate research study  |  ')
    run.font.name = 'Calibri'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string('52616B')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    doc.save(BUILD / 'reference.docx')


def make_docx(stem):
    subprocess.run([
        'pandoc', str(ROOT / (stem + '.md')), '--from=markdown-implicit_figures',
        '--to=docx', '--reference-doc=' + str(BUILD / 'reference.docx'),
        '--resource-path=' + str(ROOT),
        '-o', str(ROOT / (stem + '.docx')),
    ], check=True)
    doc = Document(ROOT / (stem + '.docx'))
    for paragraph in doc.paragraphs:
        if re.match(r'^(Figure|Table) \d+\.', paragraph.text):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        if paragraph._p.xpath('.//w:drawing'):
            paragraph.paragraph_format.keep_together = True
    for table_index, table in enumerate(doc.tables, start=1):
        table.autofit = False
        n = len(table.columns)
        total = 6.66
        ratios = TABLE_RATIOS.get(table_index, [1/n]*n)
        for i, col in enumerate(table.columns):
            col.width = Inches(total * ratios[i])
        for ri, row in enumerate(table.rows):
            props = row._tr.get_or_add_trPr()
            cant = OxmlElement('w:cantSplit')
            props.append(cant)
            if ri == 0:
                repeat = OxmlElement('w:tblHeader')
                props.append(repeat)
            for ci, cell in enumerate(row.cells):
                cell.width = Inches(total * ratios[ci])
                if ri == 0:
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:fill'), '18354D')
                    cell._tc.get_or_add_tcPr().append(shd)
                elif ri % 2:
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:fill'), 'F0F4F7')
                    cell._tc.get_or_add_tcPr().append(shd)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(4)
                    para.paragraph_format.space_before = Pt(3)
                    para.paragraph_format.line_spacing = 1.05
                    for run in para.runs:
                        run.font.name = 'Calibri'
                        run.font.size = Pt(9)
                        if ri == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
    doc.core_properties.author = ''
    doc.core_properties.subject = 'Controlled synthetic evaluation and software testing of EvidenceGate'
    doc.save(ROOT / (stem + '.docx'))


font_dir = Path('/usr/share/fonts/truetype/dejavu')
for name, fname in [
    ('ReviewSerif', 'DejaVuSerif.ttf'), ('ReviewSerif-Bold', 'DejaVuSerif-Bold.ttf'),
    ('ReviewSerif-Italic', 'DejaVuSerif-Italic.ttf'),
    ('ReviewSerif-BoldItalic', 'DejaVuSerif-BoldItalic.ttf'),
    ('ReviewSans', 'DejaVuSans.ttf'), ('ReviewSans-Bold', 'DejaVuSans-Bold.ttf'),
    ('ReviewSans-Italic', 'DejaVuSans-Oblique.ttf'),
    ('ReviewSans-BoldItalic', 'DejaVuSans-BoldOblique.ttf'),
    ('ReviewMono', 'DejaVuSansMono.ttf'),
]:
    pdfmetrics.registerFont(TTFont(name, str(font_dir / fname)))
pdfmetrics.registerFontFamily('ReviewSerif', normal='ReviewSerif', bold='ReviewSerif-Bold', italic='ReviewSerif-Italic', boldItalic='ReviewSerif-BoldItalic')
pdfmetrics.registerFontFamily('ReviewSans', normal='ReviewSans', bold='ReviewSans-Bold', italic='ReviewSans-Italic', boldItalic='ReviewSans-BoldItalic')

BODY = ParagraphStyle('Body', fontName='ReviewSerif', fontSize=10,
                      leading=14.3, spaceAfter=6, allowWidows=0,
                      allowOrphans=0, splitLongWords=True)
TITLE = ParagraphStyle('Title', fontName='ReviewSans-Bold', fontSize=23,
                       leading=29, textColor=colors.HexColor(BLUE), spaceAfter=11)
SUB = ParagraphStyle('Sub', fontName='ReviewSans', fontSize=11.5, leading=16,
                     textColor=colors.HexColor(GRAY), spaceAfter=19)
HEAD = ParagraphStyle('Head', fontName='ReviewSans-Bold', fontSize=13.2,
                      leading=17, textColor=colors.HexColor(BLUE),
                      spaceBefore=13, spaceAfter=7, keepWithNext=True)
SUBHEAD = ParagraphStyle('Subhead', parent=HEAD, fontSize=11, leading=14,
                         spaceBefore=11)
CELL = ParagraphStyle('Cell', fontName='ReviewSans', fontSize=8,
                      leading=11, spaceAfter=0, splitLongWords=True)
CELLHEAD = ParagraphStyle('CellHead', parent=CELL, fontName='ReviewSans-Bold',
                          textColor=colors.white)
CODE = ParagraphStyle('Code', fontName='ReviewMono', fontSize=8.1, leading=11,
                      leftIndent=8, rightIndent=8, spaceAfter=8)
REF = ParagraphStyle('Ref', parent=BODY, fontSize=8.4, leading=11.7,
                     spaceAfter=8)


def esc(x):
    return html.escape(x, quote=False)


def inline(nodes):
    parts = []
    for node in nodes:
        typ, c = node['t'], node.get('c')
        if typ == 'Str': parts.append(esc(c))
        elif typ in ['Space', 'SoftBreak']: parts.append(' ')
        elif typ == 'LineBreak': parts.append('<br/>')
        elif typ == 'Strong': parts.append('<b>' + inline(c) + '</b>')
        elif typ == 'Emph': parts.append('<i>' + inline(c) + '</i>')
        elif typ == 'Code': parts.append('<font name="ReviewMono" size="8.2">' + esc(c[1]) + '</font>')
        elif typ == 'Link':
            label = inline(c[1]); url = html.escape(c[2][0], quote=True)
            if len(label) > 70:
                label = label.replace('/', '/\u200b').replace('_', '_\u200b')
            parts.append(f'<a href="{url}" color="#225780">{label}</a>')
        elif typ == 'Quoted':
            parts.append('“' + inline(c[1]) + '”')
        elif typ == 'Span': parts.append(inline(c[1]))
        elif typ == 'Math': parts.append(esc(c[1]))
        elif typ == 'Cite': parts.append(inline(c[1]))
        elif typ == 'RawInline': pass
        else: raise ValueError('Unhandled inline: ' + typ)
    return ''.join(parts)


def meta_text(value):
    if value['t'] == 'MetaInlines': return inline(value['c'])
    return ''


def make_pdf(stem):
    ast = json.loads(subprocess.check_output([
        'pandoc', str(ROOT/(stem+'.md')), '-f', 'markdown-implicit_figures', '-t', 'json'
    ]))
    out = ROOT/(stem+'.pdf')
    width = A4[0] - 96
    doc = SimpleDocTemplate(str(out), pagesize=A4, topMargin=49,
                            bottomMargin=48, leftMargin=48, rightMargin=48,
                            title=html.unescape(re.sub('<[^>]+>', '', meta_text(ast['meta']['title']))) if 'title' in ast['meta'] else stem.replace('_', ' '),
                            author='', pageCompression=1)
    story = []
    if 'title' in ast['meta']:
        story.append(Paragraph(meta_text(ast['meta']['title']), TITLE))
    if 'subtitle' in ast['meta']:
        story.append(Paragraph(meta_text(ast['meta']['subtitle']), SUB))
    in_refs = False
    table_index = 0

    def block_list(blocks, listlevel=0):
        nonlocal in_refs, table_index
        result=[]
        for block in blocks:
            typ, c = block['t'], block.get('c')
            if typ == 'Header':
                text=inline(c[2]); level=c[0]
                if text == 'References': in_refs = True
                sty = TITLE if level == 1 and not ast['meta'] else (HEAD if level <= 2 else SUBHEAD)
                result.append(Paragraph(text, sty))
            elif typ in ['Para', 'Plain']:
                image_nodes = [node for node in c if node['t'] == 'Image']
                if image_nodes:
                    for node in image_nodes:
                        filename = ROOT / node['c'][2][0]
                        drawing = ReportlabImage(str(filename))
                        ratio = drawing.imageHeight / drawing.imageWidth
                        drawing.drawWidth = min(width, 6.5 * 72)
                        drawing.drawHeight = drawing.drawWidth * ratio
                        drawing.hAlign = 'CENTER'
                        result.extend([drawing, Spacer(1, 8)])
                    continue
                text=inline(c)
                style=REF if in_refs else BODY
                if text.startswith('<b>Table ') or text.startswith('<b>Figure '):
                    style=ParagraphStyle('Caption', parent=BODY, fontName='ReviewSans', fontSize=9,
                                         leading=12, keepWithNext=True, spaceBefore=7)
                result.append(Paragraph(text, style))
            elif typ in ['BulletList', 'OrderedList']:
                items = c if typ=='BulletList' else c[1]
                for i, item in enumerate(items):
                    for j, child in enumerate(item):
                        if child['t'] in ['Para', 'Plain']:
                            style=ParagraphStyle('List', parent=BODY, leftIndent=14,
                                                 firstLineIndent=-10 if j==0 else 0)
                            prefix = ('• ' if typ=='BulletList' else str(i+1)+'. ') if j==0 else ''
                            result.append(Paragraph(prefix+inline(child['c']),style))
                        else: result.extend(block_list([child],listlevel+1))
            elif typ == 'CodeBlock':
                code = c[1]
                result.append(Preformatted(code,CODE,maxLineLength=83))
            elif typ == 'Table':
                table_index += 1
                # Pandoc table layout: attr, caption, colspec, head, bodies, foot.
                rows=[(r,True) for r in c[3][1]]
                for body in c[4]:
                    rows.extend((r,False) for r in body[2])
                    rows.extend((r,False) for r in body[3])
                data=[]
                for row, header in rows:
                    values=[]
                    for cell in row[1]:
                        paras=[]
                        for b in cell[4]:
                            if b['t'] in ['Para','Plain']: paras.append(inline(b['c']))
                            else: raise ValueError('Unexpected table block '+b['t'])
                        values.append(Paragraph('<br/>'.join(paras), CELLHEAD if header else CELL))
                    data.append(values)
                n=len(data[0]); ratios=TABLE_RATIOS.get(table_index,[1/n]*n)
                table=Table(data,colWidths=[width*r for r in ratios],repeatRows=1,hAlign='LEFT')
                table.setStyle(TableStyle([
                    ('BACKGROUND',(0,0),(-1,0),colors.HexColor(BLUE)),
                    ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#f0f4f7'),colors.white]),
                    ('VALIGN',(0,0),(-1,-1),'TOP'),
                    ('LEFTPADDING',(0,0),(-1,-1),7),
                    ('RIGHTPADDING',(0,0),(-1,-1),7),
                    ('TOPPADDING',(0,0),(-1,-1),7),
                    ('BOTTOMPADDING',(0,0),(-1,-1),7),
                    ('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor(BLUE)),
                    ('LINEBELOW',(0,-1),(-1,-1),.5,colors.HexColor('#bacad6')),
                ]))
                result.extend([table,Spacer(1,9)])
            elif typ == 'HorizontalRule': result.append(Spacer(1,9))
            elif typ == 'Div': result.extend(block_list(c[1]))
            else: raise ValueError('Unhandled block '+typ)
        return result

    conclusion_index = next((i for i, block in enumerate(ast['blocks'])
                             if block['t'] == 'Header'
                             and re.match(r'^\d+\. Conclusion$', inline(block['c'][2]))), None)
    if conclusion_index is None:
        story.extend(block_list(ast['blocks']))
    else:
        story.extend(block_list(ast['blocks'][:conclusion_index]))
        story.append(KeepTogether(block_list(ast['blocks'][conclusion_index:])))

    def page(canvas, document):
        canvas.saveState()
        canvas.setFont('ReviewSans',7.4)
        canvas.setFillColor(colors.HexColor(GRAY))
        label = 'EvidenceGate | Controlled synthetic evaluation'
        canvas.drawString(48,27,label)
        canvas.drawRightString(A4[0]-48,27,str(document.page))
        canvas.setStrokeColor(colors.HexColor('#d4dfe5'))
        canvas.line(48,39,A4[0]-48,39)
        canvas.restoreState()
    doc.build(story,onFirstPage=page,onLaterPages=page)


if __name__ == '__main__':
    make_reference()
    make_docx('EvidenceGate_Research_Paper')
    make_pdf('EvidenceGate_Research_Paper')
    print('Built editable Word document and PDF manuscript.')
